"""
Test script to verify the document processing functionality.

This script tests the core functions without needing the web interface.
"""

import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "Application"))

from app import extract_placeholders, fill_document
from docx import Document


def test_placeholder_extraction():
    """Test placeholder extraction from SAFE document."""
    print("=" * 70)
    print("Testing Placeholder Extraction")
    print("=" * 70)

    doc_path = os.path.join(os.path.dirname(__file__), "test_safe.docx")
    placeholders = extract_placeholders(doc_path)

    print(f"\n Found {len(placeholders)} placeholders:\n")

    for i, placeholder in enumerate(placeholders, 1):
        print(f"{i}. {placeholder['original']}")
        print(f"   Name: {placeholder['name']}")
        print(f"   Type: {placeholder['type']}")
        print(f"   Context: ...{placeholder['context'][:60]}...")
        print()

    return placeholders


def test_document_filling():
    """Test document filling with sample data."""
    print("\n" + "=" * 70)
    print("Testing Document Filling")
    print("=" * 70)

    # First, extract placeholders
    template_path = os.path.join(os.path.dirname(__file__), "test_safe.docx")
    placeholders = extract_placeholders(template_path)

    # Create responses dict that matches what the app expects
    # Format: {index: {placeholder: info, value: answer}}
    responses = {}

    # Sample answers for the placeholders
    sample_answers = {
        "Company Name": "ACM Technologies Inc.",
        "Investor Name": "Naveen Alampally",
        "Purchase Amount": "$1,000,000",
        "Date of Safe": "January 15, 2025",
        "State of Incorporation": "Texas",
        "Post-Money Valuation Cap": "$5,000,000",
        "Governing Law Jurisdiction": "Texas",
        "COMPANY": "ACM TECHNOLOGIES INC.",
        "name": "Naveen Alampally",
        "title": "CEO",
    }

    print("\n Filling document with sample data:")

    # Map placeholders to responses
    for idx, placeholder in enumerate(placeholders):
        name = placeholder["name"]
        if name in sample_answers:
            responses[str(idx)] = {
                "placeholder": placeholder,
                "value": sample_answers[name],
            }
            print(f"   {placeholder['original']} → {sample_answers[name]}")

    # Fill the document
    output_path = os.path.join(os.path.dirname(__file__), "test_output.docx")

    fill_document(template_path, output_path, responses, placeholders)

    print(f"\n Document filled successfully!")
    print(f"   Output: {output_path}")

    # Verify the output
    doc = Document(output_path)
    full_text = "\n".join([para.text for para in doc.paragraphs])

    print("\n🔍 Verification:")
    for name, value in sample_answers.items():
        if value in full_text:
            print(f"    {name} filled with '{value}'")

    return output_path


def main():
    """Run all tests."""
    print("\n Legal Document Filler - Test Suite\n")

    try:
        # Test 1: Extract placeholders
        placeholders = test_placeholder_extraction()

        # Test 2: Fill document
        output_path = test_document_filling()

        print("\n" + "=" * 70)
        print(" All Tests Passed!")
        print("=" * 70)
        print(f"\nTest output document: {output_path}")
        print("You can open this document to verify the placeholders were filled.\n")

    except Exception as e:
        print(f"\n Test failed: {e}")
        import traceback

        traceback.print_exc()


if __name__ == "__main__":
    main()
